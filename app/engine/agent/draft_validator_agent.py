"""
draft_validator_agent.py  (v2)

Validates the *final rendered .docx bytes* produced by TemplateEngine +
PostProcessor — not raw text pages.

Pipeline position:
    PostProcessor.process_package()  →  DraftValidatorAgent.validate()

The validator opens each .docx with python-docx, extracts full paragraph text,
then runs the same four category checks as v1, now against real rendered output:

    1. structure      — mandatory doc keys present in the rendered set
    2. content_length — each document has a minimum paragraph / word count
    3. completeness   — required markers found in the primary document's text
    4. legal_format   — draft-type section headings present in primary doc

VALIDATION_WEIGHTS / COMPLETENESS_MARKERS / LEGAL_FORMAT_CHECKS are imported
from draft_constants (unchanged from v1) so tuning stays in one place.
"""

from __future__ import annotations

import io
import re
from typing import Dict, List, Optional, Tuple

from docx import Document
from loguru import logger

from app.engine.draft_constants import (
    COMPLETENESS_MARKERS,
    DRAFT_DOCUMENTS,
    LEGAL_FORMAT_CHECKS,
    VALIDATION_PASS_THRESHOLD,
    VALIDATION_WEIGHTS,
    ValidationResult,
    ValidationStatus,
)

# ── Minimum content thresholds ────────────────────────────────────────────────
# A rendered document that falls below these is almost certainly empty/broken.
_MIN_WORDS_PER_DOC = 50  # across all paragraphs in the doc
_MIN_PARAGRAPHS_PER_DOC = 5  # non-empty paragraphs


def _slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^\w\s-]", "", text)
    text = re.sub(r"[\s_-]+", "_", text)
    return text


def _extract_text(docx_bytes: bytes) -> str:
    """Return all paragraph text from a .docx file, joined by spaces."""
    doc = Document(io.BytesIO(docx_bytes))
    return " ".join(p.text.strip() for p in doc.paragraphs if p.text.strip())


def _extract_paragraphs(docx_bytes: bytes) -> List[str]:
    """Return list of non-empty paragraph strings from a .docx file."""
    doc = Document(io.BytesIO(docx_bytes))
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


class DraftValidatorAgent:
    """
    Validates rendered .docx bytes.

    Usage:
        validator = DraftValidatorAgent()
        result = validator.validate(draft_type, rendered_docs)

    Args:
        draft_type:    e.g. "writ_petition", "bail_application"
        rendered_docs: {doc_key: docx_bytes}  — output of PostProcessor
    """

    def validate(
        self,
        draft_type: str,
        rendered_docs: Dict[str, bytes],
    ) -> ValidationResult:

        all_issues: List[str] = []
        all_passed: List[str] = []
        cat_scores: Dict[str, float] = {}

        # ── 1. Structure: mandatory doc keys present ───────────────────────────
        cat_scores["structure"], s_passed, s_issues = self._check_structure(
            draft_type, rendered_docs
        )
        all_passed.extend(s_passed)
        all_issues.extend(s_issues)

        # ── 2. Content length: each doc has enough text ────────────────────────
        cat_scores["content_length"], l_passed, l_issues = self._check_content_length(
            rendered_docs
        )
        all_passed.extend(l_passed)
        all_issues.extend(l_issues)

        # ── 3. Completeness: markers in primary document ───────────────────────
        main_key, main_text = self._get_main_text(draft_type, rendered_docs)
        cat_scores["completeness"], c_passed, c_issues = self._check_completeness(
            main_key, main_text
        )
        all_passed.extend(c_passed)
        all_issues.extend(c_issues)

        # ── 4. Legal format: section headings in primary document ──────────────
        cat_scores["legal_format"], f_passed, f_issues = self._check_legal_format(
            draft_type, main_key, main_text
        )
        all_passed.extend(f_passed)
        all_issues.extend(f_issues)

        # ── Final score ────────────────────────────────────────────────────────
        # VALIDATION_WEIGHTS must have keys matching cat_scores keys.
        # If a key is missing we skip it gracefully.
        final_score = round(
            sum(
                VALIDATION_WEIGHTS.get(cat, 0.0) * score
                for cat, score in cat_scores.items()
            ),
            4,
        )
        final_score = min(final_score, 1.0)

        status = (
            ValidationStatus.VALID
            if final_score >= VALIDATION_PASS_THRESHOLD
            else ValidationStatus.NEEDS_REVISION
        )

        logger.info(
            f"DraftValidatorAgent>>result "
            f"draft_type={draft_type} score={final_score} status={status} "
            f"passed={len(all_passed)} issues={len(all_issues)} "
            f"category_scores={cat_scores}"
        )

        return ValidationResult(
            score=final_score,
            status=status,
            issues=all_issues,
            passed_checks=all_passed,
        )

    # ── Private checks ────────────────────────────────────────────────────────

    def _check_structure(
        self,
        draft_type: str,
        rendered_docs: Dict[str, bytes],
    ) -> Tuple[float, List[str], List[str]]:
        """
        Every mandatory doc key for this draft_type must be present in
        rendered_docs.  Slug-based fuzzy matching is kept for resilience.
        """
        registry = DRAFT_DOCUMENTS.get(draft_type, {"mandatory": [], "optional": []})
        mandatory = registry.get("mandatory", [])

        slug_map = {_slugify(k): k for k in rendered_docs.keys()}
        passed, issues = [], []

        for doc_key in mandatory:
            doc_slug = _slugify(doc_key)
            matched = any(doc_slug in actual_slug for actual_slug in slug_map)
            if matched:
                passed.append(f"mandatory doc rendered: '{doc_key}'")
            else:
                issues.append(f"missing mandatory rendered document: '{doc_key}'")

        score = len(passed) / len(mandatory) if mandatory else 1.0
        return score, passed, issues

    def _check_content_length(
        self,
        rendered_docs: Dict[str, bytes],
    ) -> Tuple[float, List[str], List[str]]:
        """
        Each rendered .docx must meet minimum paragraph and word-count thresholds.
        Catches blank/corrupted templates before they reach the client.
        """
        passed, issues = [], []

        for doc_key, docx_bytes in rendered_docs.items():
            try:
                paragraphs = _extract_paragraphs(docx_bytes)
                word_count = sum(len(p.split()) for p in paragraphs)

                if len(paragraphs) < _MIN_PARAGRAPHS_PER_DOC:
                    issues.append(
                        f"content_length: '{doc_key}' has only {len(paragraphs)} "
                        f"non-empty paragraphs (min {_MIN_PARAGRAPHS_PER_DOC})"
                    )
                elif word_count < _MIN_WORDS_PER_DOC:
                    issues.append(
                        f"content_length: '{doc_key}' has only {word_count} words "
                        f"(min {_MIN_WORDS_PER_DOC})"
                    )
                else:
                    passed.append(
                        f"content_length: '{doc_key}' ok "
                        f"({len(paragraphs)} paragraphs, {word_count} words)"
                    )
            except Exception as exc:
                issues.append(f"content_length: '{doc_key}' could not be read — {exc}")

        total = len(rendered_docs)
        score = len(passed) / total if total else 1.0
        return score, passed, issues

    def _get_main_text(
        self,
        draft_type: str,
        rendered_docs: Dict[str, bytes],
    ) -> Tuple[Optional[str], str]:
        """
        Return (main_doc_key, full_text) for the primary document.
        Primary = first mandatory doc; falls back to first key.
        """
        registry = DRAFT_DOCUMENTS.get(draft_type, {"mandatory": [], "optional": []})
        mandatory = registry.get("mandatory", [])

        main_key: Optional[str] = None

        if mandatory:
            target_slug = _slugify(mandatory[0])
            for actual_key in rendered_docs.keys():
                if target_slug in _slugify(actual_key):
                    main_key = actual_key
                    break

        if main_key is None and rendered_docs:
            main_key = list(rendered_docs.keys())[0]

        if main_key and main_key in rendered_docs:
            try:
                text = _extract_text(rendered_docs[main_key])
                return main_key, text
            except Exception as exc:
                logger.warning(
                    f"DraftValidatorAgent>>could not extract text from "
                    f"'{main_key}': {exc}"
                )

        return main_key, ""

    def _check_completeness(
        self,
        main_key: Optional[str],
        main_text: str,
    ) -> Tuple[float, List[str], List[str]]:
        """
        Search the primary document's full text for COMPLETENESS_MARKERS.
        Markers are (label, regex_pattern) tuples from draft_constants.
        """
        passed, issues = [], []

        for label, pattern in COMPLETENESS_MARKERS:
            if re.search(pattern, main_text, re.IGNORECASE):
                passed.append(f"completeness: {label}")
            else:
                issues.append(f"completeness: '{label}' not found in '{main_key}'")

        score = len(passed) / len(COMPLETENESS_MARKERS) if COMPLETENESS_MARKERS else 1.0
        return score, passed, issues

    def _check_legal_format(
        self,
        draft_type: str,
        main_key: Optional[str],
        main_text: str,
    ) -> Tuple[float, List[str], List[str]]:
        """
        Run draft-type-specific heading / section checks against main_text.
        Rules are (label, regex_pattern) tuples from LEGAL_FORMAT_CHECKS.
        """
        format_rules = LEGAL_FORMAT_CHECKS.get(draft_type, [])
        passed, issues = [], []

        for check_label, pattern in format_rules:
            if re.search(pattern, main_text, re.IGNORECASE):
                passed.append(f"legal format: {check_label}")
            else:
                issues.append(
                    f"legal format: '{check_label}' — section not found in '{main_key}'"
                )

        score = len(passed) / len(format_rules) if format_rules else 1.0
        return score, passed, issues
